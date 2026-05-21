from __future__ import annotations

from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.db.init_db import seed_demo_data
from app.main import app


seed_demo_data()
client = TestClient(app)


def auth_headers() -> dict:
    login = client.post(
        "/api/v1/auth/login",
        json={"email": "daniel@example.com", "password": "securepassword"},
    )
    assert login.status_code == 200, login.text
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


def test_report_generation_creates_local_pdf_and_excel() -> None:
    headers = auth_headers()
    pdf = client.post(
        "/api/v1/reports/generate",
        headers=headers,
        json={"report_type": "risk", "format": "pdf"},
    )
    assert pdf.status_code == 200, pdf.text
    pdf_path = Path("generated_reports") / Path(pdf.json()["download_url"]).name
    assert pdf_path.exists()
    assert pdf_path.stat().st_size > 100

    excel = client.post(
        "/api/v1/reports/generate",
        headers=headers,
        json={"report_type": "forecast", "format": "xlsx"},
    )
    assert excel.status_code == 200, excel.text
    excel_path = Path("generated_reports") / Path(excel.json()["download_url"]).name
    assert excel_path.exists()
    workbook = load_workbook(excel_path)
    assert {"Summary", "Risk Breakdown", "Forecast", "Recommendations", "Scenario"}.issubset(
        set(workbook.sheetnames)
    )


def test_protected_routes_require_authentication() -> None:
    response = client.get("/api/v1/business/me")
    assert response.status_code == 401


def test_report_templates_include_required_types() -> None:
    response = client.get("/api/v1/reports/templates")
    assert response.status_code == 200
    keys = {item["key"] for item in response.json()["templates"]}
    assert {"risk", "forecast", "scenario", "lender"}.issubset(keys)
