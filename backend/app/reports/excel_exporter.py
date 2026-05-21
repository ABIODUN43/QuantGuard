from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


GREEN = "009B63"
LIGHT_GREEN = "ECFDF5"
HEADER = "0F172A"


def _write_rows(ws, rows: list[tuple], start_row: int = 1) -> None:
    for row_index, row in enumerate(rows, start=start_row):
        for column_index, value in enumerate(row, start=1):
            cell = ws.cell(row=row_index, column=column_index, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_index == start_row:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=GREEN)


def _autosize(ws) -> None:
    for column_cells in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_length + 2, 14), 42)


def _summary_rows(payload: dict) -> list[tuple]:
    business = payload["business"]
    analysis = payload["analysis"]
    forecast = payload["forecast"]
    return [
        ("Metric", "Value"),
        ("Business", business["business_name"]),
        ("Industry", business["industry"]),
        ("Location", f"{business['city']}, {business['country']}"),
        ("Risk Score", analysis["risk_score"]),
        ("Risk Level", analysis["risk_level"]),
        ("Stress Probability", analysis["stress_probability"]),
        ("Forecast Confidence", analysis["forecast_confidence"]),
        ("Projected Cash Balance", forecast["projected_cash_balance"]),
        ("Cash Shortfall Probability", forecast["cash_shortfall_probability"]),
        ("Danger Period", forecast["danger_period"]),
        ("Breakeven Revenue", forecast["breakeven_revenue"]),
    ]


def generate_excel_report(payload: dict, output_path: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "QuantGuard AI"
    summary["A1"].font = Font(bold=True, color=GREEN, size=14)
    summary["A2"] = payload["template"]["title"]
    summary["A2"].font = Font(bold=True, color=HEADER, size=18)
    _write_rows(summary, _summary_rows(payload), start_row=4)
    _autosize(summary)

    risk = workbook.create_sheet("Risk Breakdown")
    risk_rows = [("Risk Factor", "Score", "Contribution %")]
    contributions = payload["analysis"].get("contribution_percentages", {})
    for key, value in payload["analysis"]["risk_breakdown"].items():
        risk_rows.append((key.replace("_", " ").title(), value, contributions.get(key, 0)))
    _write_rows(risk, risk_rows)
    _autosize(risk)

    forecast = workbook.create_sheet("Forecast")
    forecast_rows = [("Month", "Revenue", "Expenses", "Projected Cash", "Low Cash", "High Cash")]
    for point in payload["forecast"].get("cashflow_forecast", []):
        forecast_rows.append(
            (
                point["month"],
                point["projected_revenue"],
                point["projected_expenses"],
                point["projected_cash_balance"],
                point.get("low_cash_balance"),
                point.get("high_cash_balance"),
            )
        )
    _write_rows(forecast, forecast_rows)
    _autosize(forecast)

    recommendations = workbook.create_sheet("Recommendations")
    recommendation_rows = [("Title", "Description", "Priority", "Impact", "Category")]
    for item in payload.get("recommendations", []):
        recommendation_rows.append(
            (item["title"], item["description"], item["priority"], item["impact"], item["category"])
        )
    _write_rows(recommendations, recommendation_rows)
    _autosize(recommendations)

    scenario = workbook.create_sheet("Scenario")
    scenario_data = payload.get("scenario")
    if scenario_data:
        _write_rows(
            scenario,
            [
                ("Metric", "Value"),
                ("Scenario", scenario_data["scenario_name"]),
                ("Risk Score", scenario_data["risk_score"]),
                ("Stress Probability", scenario_data["stress_probability"]),
                ("Cash Shortfall Probability", scenario_data["cash_shortfall_probability"]),
                ("Impact", scenario_data["impact"]),
                ("Summary", scenario_data["summary"]),
            ],
        )
    else:
        _write_rows(scenario, [("Metric", "Value"), ("Scenario", "No scenario generated yet")])
    _autosize(scenario)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows():
            for cell in row:
                if cell.row > 1 and cell.fill.fill_type is None:
                    cell.fill = PatternFill("solid", fgColor=LIGHT_GREEN if cell.row % 2 == 0 else "FFFFFF")

    workbook.save(output_path)
